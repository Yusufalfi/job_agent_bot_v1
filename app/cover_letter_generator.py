import os
import json
import re
from datetime import datetime

import httpx
from docx import Document
from groq import Groq


def load_profile(json_path=None) -> dict:
    """
    Membaca profile.json dari folder config.
    Struktur project:

    Agentic_Ai_Job_match/
    ├── app/
    │   └── coverLeter.py
    └── config/
        └── profile.json
    """

    # Kalau path tidak diberikan,
    # cari project root berdasarkan lokasi file ini.
    if json_path is None:
        project_root = os.path.dirname(
            os.path.dirname(os.path.abspath(__file__))
        )

        json_path = os.path.join(
            project_root,
            "config",
            "profile.json"
        )

    if not os.path.exists(json_path):
        raise FileNotFoundError(
            f"profile.json tidak ditemukan: {json_path}"
        )

    try:
        with open(json_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        if not isinstance(data, dict):
            raise ValueError("Isi profile.json harus berupa object JSON.")

        return data

    except json.JSONDecodeError as e:
        raise ValueError(
            f"Format profile.json tidak valid: {e}"
        )


def generate_cover_letter_docx(
    job: dict,
    profile_path: str = None
) -> str:

    # ==============================
    # 1. Ambil API Key Groq
    # ==============================

    api_key = os.getenv("GROQ_API_KEY")

    if not api_key:
        raise ValueError(
            "GROQ_API_KEY tidak ditemukan! "
            "Pastikan file .env sudah terisi dengan benar."
        )

    # ==============================
    # 2. Load Profile
    # ==============================

    profile = load_profile(profile_path)

    # Ambil data personal dari profile.json
    name = profile.get("name")
    address = profile.get("address")
    email = profile.get("email")
    phone = profile.get("phone")

    # Pastikan data personal tersedia
    if not all([name, address, email, phone]):
        raise ValueError(
            "Data personal di profile.json belum lengkap. "
            "Pastikan name, address, email, dan phone tersedia."
        )

    # ==============================
    # 3. Ambil Skills
    # ==============================

    skills_list = profile.get("keywords", [])

    skills_str = (
        ", ".join(skills_list)
        if skills_list
        else "keahlian yang relevan"
    )

    # ==============================
    # 4. Buat Groq Client
    # ==============================

    client = Groq(
        api_key=api_key,
        http_client=httpx.Client(verify=False)
    )

    # ==============================
    # 5. Buat Prompt
    # ==============================

    prompt = f"""
Buatkan draf isi Cover Letter / Surat Lamaran kerja
yang profesional, ringkas, dan persuasif dalam Bahasa Indonesia.

Surat ini ditujukan untuk melamar pekerjaan berikut:

- Posisi: {job.get('title', 'Tidak diketahui')}
- Perusahaan: {job.get('company', 'Tidak diketahui')}
- Deskripsi Pekerjaan:
{str(job.get('description', ''))[:4000]}

Profil Pelamar:

- Target Roles: {profile.get('target_roles', [])}
- Keahlian Utama (Skills): {skills_str}

Aturan:

1. Jangan mengarang pengalaman atau kemampuan yang tidak tercantum
   dalam profil pelamar.

2. Fokuskan penekanan pada kesesuaian skill pelamar
   ({skills_str}) dengan kebutuhan pekerjaan di atas.

3. Surat harus ditulis dari sudut pandang pelamar.
   Pelamar adalah orang yang sedang melamar pekerjaan,
   bukan perusahaan atau job portal.

4. HANYA berikan isi paragraf surat lamarannya saja.

5. Jangan sertakan:
   - Nama pelamar
   - Alamat
   - Email
   - Nomor telepon
   - Header surat
   - "Kepada Yth..."
   - Perihal
   - Tanda tangan

6. Mulai langsung dari isi paragraf pembuka
   setelah "Dengan hormat,".

7. Jangan menggunakan placeholder seperti:
   "Nama Anda", "Nama Perusahaan", "Alamat", dan sebagainya.
"""

    # ==============================
    # 6. Request ke Groq
    # ==============================

    response = client.chat.completions.create(
        model="openai/gpt-oss-120b",
        messages=[
            {
                "role": "user",
                "content": prompt
            }
        ],
        temperature=0.3,
        max_completion_tokens=1024
    )

    letter_content = response.choices[0].message.content

    if not letter_content:
        raise ValueError(
            "AI tidak menghasilkan isi Cover Letter."
        )

    # ==============================
    # 7. Membuat Dokumen Word
    # ==============================

    doc = Document()

    # Header Data Diri
    doc.add_paragraph(
        f"{name}\n"
        f"{address}\n"
        f"{email} | {phone}"
    )

    # Tanggal
    today_str = datetime.now().strftime("%d %B %Y")

    doc.add_paragraph(
        f"\nBekasi, {today_str}\n"
    )

    # Tujuan Surat
    company = job.get(
        "company",
        "Perusahaan"
    )

    title = job.get(
        "title",
        "Job"
    )

    doc.add_paragraph(
        f"Kepada Yth.\n"
        f"Manajer Rekrutmen {company}\n"
        f"Di Tempat\n"
    )

    # Perihal
    doc.add_heading(
        f"Perihal: Lamaran untuk Posisi {title}",
        level=2
    )

    # Salam
    doc.add_paragraph(
        "Dengan hormat,"
    )

    # ==============================
    # 8. Masukkan Isi dari AI
    # ==============================

    for paragraph in letter_content.split("\n\n"):

        if paragraph.strip():
            doc.add_paragraph(
                paragraph.strip()
            )

    # ==============================
    # 9. Penutup
    # ==============================

    doc.add_paragraph(
        "\nDemikian surat lamaran ini saya sampaikan. "
        "Atas perhatian dan kesempatan yang diberikan, "
        "saya ucapkan terima kasih."
    )

    doc.add_paragraph(
        f"\nHormat saya,\n\n\n{name}"
    )

    # ==============================
    # 10. Simpan ke Folder data
    # ==============================

    project_root = os.path.dirname(
        os.path.dirname(os.path.abspath(__file__))
    )

    output_dir = os.path.join(
        project_root,
        "data"
    )

    os.makedirs(
        output_dir,
        exist_ok=True
    )

    # Bersihkan nama perusahaan
    company_clean = re.sub(
        r'[^\w\s-]',
        '',
        company
    ).strip()

    filename = (
        f"Cover_Letter_"
        f"{company_clean.replace(' ', '_')}.docx"
    )

    full_path = os.path.join(
        output_dir,
        filename
    )

    # Simpan Word
    doc.save(full_path)

    return full_path