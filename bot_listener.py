import os
import time
import json
import logging
import requests
from dotenv import load_dotenv

# Import generator Cover Letter yang sudah kamu buat
from app.cover_letter_generator import generate_cover_letter_docx
from app.database import EXCEL_FILE_PATH
import openpyxl

load_dotenv()
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
CHAT_ID = os.getenv("TELEGRAM_CHAT_ID")

if not TOKEN:
    print("❌ TELEGRAM_BOT_TOKEN belum disetel di file .env!")
    exit(1)

def get_job_from_excel(target_job_id: str) -> dict:
    if not os.path.exists(EXCEL_FILE_PATH):
        return None
    
    wb = openpyxl.load_workbook(EXCEL_FILE_PATH, data_only=True)
    ws = wb.active
    
    for row in range(2, ws.max_row + 1):
        job_id = str(ws.cell(row=row, column=1).value)
        if job_id == str(target_job_id):
            return {
                "job_id": job_id,
                "title": ws.cell(row=row, column=3).value,
                "company": ws.cell(row=row, column=4).value,
                "location": ws.cell(row=row, column=5).value,
                "salary_text": ws.cell(row=row, column=6).value,
                "description": ws.cell(row=row, column=9).value,
                "work_type": ws.cell(row=row, column=10).value,
            }
    wb.close()
    return None

def answer_callback_query(callback_query_id: str):
    url = f"https://api.telegram.org/bot{TOKEN}/answerCallbackQuery"
    try:
        requests.post(url, json={"callback_query_id": callback_query_id}, verify=False, timeout=5)
    except Exception:
        pass

# Diperbarui agar bisa menerima reply_markup untuk menghapus tombol
def edit_message_text(chat_id, message_id, new_text, remove_keyboard=False):
    url = f"https://api.telegram.org/bot{TOKEN}/editMessageText"
    payload = {
        "chat_id": chat_id,
        "message_id": message_id,
        "text": new_text,
        "parse_mode": "Markdown"
    }
    if remove_keyboard:
        payload["reply_markup"] = {"inline_keyboard": []}

    try:
        requests.post(url, json=payload, verify=False, timeout=10)
    except Exception as e:
        logging.error(f"Gagal edit pesan: {e}")

def send_document_to_telegram(chat_id, file_path, caption):
    url = f"https://api.telegram.org/bot{TOKEN}/sendDocument"
    try:
        with open(file_path, "rb") as doc_file:
            files = {"document": doc_file}
            data = {
                "chat_id": chat_id,
                "caption": caption,
                "parse_mode": "Markdown"
            }
            response = requests.post(url, data=data, files=files, verify=False, timeout=30)
            if response.status_code == 200:
                logging.info("File Cover Letter berhasil dikirim ke Telegram!")
            else:
                logging.error(f"Gagal kirim dokumen: {response.text}")
    except Exception as e:
        logging.error(f"Error saat upload dokumen: {e}")

def send_text_message(chat_id, text):
    url = f"https://api.telegram.org/bot{TOKEN}/sendMessage"
    payload = {
        "chat_id": chat_id,
        "text": text,
        "parse_mode": "Markdown"
    }
    try:
        requests.post(url, json=payload, verify=False, timeout=10)
    except Exception as e:
        logging.error(f"Gagal kirim pesan teks: {e}")

def main():
    print("🤖 Bot Listener (Requests Mode) aktif dan siap mendengarkan klik tombol...")
    offset = 0

    while True:
        url = f"https://api.telegram.org/bot{TOKEN}/getUpdates"
        params = {"offset": offset, "timeout": 30}
        
        try:
            response = requests.get(url, params=params, verify=False, timeout=35)
            if response.status_code == 200:
                data = response.json()
                for result in data.get("result", []):
                    offset = result["update_id"] + 1
                    
                    if "callback_query" in result:
                        callback = result["callback_query"]
                        callback_id = callback["id"]
                        callback_data = callback["data"]
                        message = callback["message"]
                        chat_id = message["chat"]["id"]
                        message_id = message["message_id"]
                        original_text = message.get("text", "")

                        # Hilangkan animasi loading di tombol HP
                        answer_callback_query(callback_id)

                        if callback_data.startswith("gen_cl_"):
                            job_id = callback_data.replace("gen_cl_", "")
                            
                            # Update teks dan HAPUS TOMBOLNYA (remove_keyboard=True)
                            updated_text = f"{original_text}\n\n⏳ *Status:* Oke, saya siapkan cover letter-nya..."
                            edit_message_text(chat_id, message_id, updated_text, remove_keyboard=True)

                            # Ambil data lowongan dari Excel
                            job_data = get_job_from_excel(job_id)

                            if not job_data:
                                send_text_message(chat_id, "❌ Maaf, data lowongan tidak ditemukan di database Excel.")
                                continue

                            try:
                                # Generate file .docx
                                file_path = generate_cover_letter_docx(job_data)

                                # Kirim dokumen Word ke Telegram
                                caption = (
                                    f"✅ **Cover Letter Selesai!**\n\n"
                                    f"🏢 *Perusahaan:* {job_data.get('company')}\n"
                                    f"💼 *Posisi:* {job_data.get('title')}\n\n"
                                    f"File sudah terisi datamu dan siap digunakan."
                                )
                                send_document_to_telegram(chat_id, file_path, caption)

                            except Exception as e:
                                logging.error(f"Gagal generate cover letter: {e}")
                                send_text_message(chat_id, f"❌ Terjadi kesalahan saat membuat Cover Letter: {str(e)}")

                        elif callback_data.startswith("skip_"):
                            # Hapus tombol juga saat di-skip
                            updated_text = f"{original_text}\n\n❌ *Status:* Lowongan dilewati."
                            edit_message_text(chat_id, message_id, updated_text, remove_keyboard=True)

            else:
                time.sleep(3)
        except requests.exceptions.RequestException as e:
            logging.warning(f"Koneksi terputus sementara, mencoba ulang... ({e})")
            time.sleep(5)
        except Exception as ex:
            logging.error(f"Terjadi error tak terduga: {ex}")
            time.sleep(3)

if __name__ == "__main__":
    main()